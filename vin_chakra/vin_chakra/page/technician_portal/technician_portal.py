import frappe
from frappe import _
from frappe.utils import add_days, flt, formatdate, get_fullname, get_user_date_format, getdate, today
from io import BytesIO
import json
import math
import re
from typing import Union

from vin_chakra.technician_api import enrich_tickets_customer_details

# Ticket columns shown by the portal's card, list, calendar and Expense Claim views
TICKET_COLUMNS = """
    name,
    subject,
    status,
    priority,
    custom_customer_name,
    custom_customer_mobile_number,
    custom__secondary_phone_number,
    custom_address,
    custom_city__district_,
    custom_state,
    custom_machine_name,
    custom_machine_problem,
    custom_required_spares,
    custom_date,
    creation,
    modified,
    _assign,
    customer,
    contact
"""

@frappe.whitelist()
def get_portal_data(
    status: str = None,
    priority: str = None,
    search_query: str = None,
    limit_start: int = 0,
    limit_page_length: int = 10,
    view: str = "tickets"
):
    user = frappe.session.user
    if user == "Guest":
        frappe.throw(_("Authentication required."), frappe.PermissionError)
        
    cache_key = f"tech_portal:{user}:{status}:{priority}:{search_query}:{limit_start}:{limit_page_length}:{view}"
    cached_res = frappe.cache().get_value(cache_key)
    if cached_res:
        return frappe.parse_json(cached_res)
        
    conditions = ["JSON_SEARCH(`_assign`, 'one', %(user)s) IS NOT NULL"]
    values = {"user": user}
    
    if status:
        conditions.append("status = %(status)s")
        values["status"] = status
    if priority:
        conditions.append("priority = %(priority)s")
        values["priority"] = priority
    if search_query:
        search_escaped = f"%{search_query}%"
        conditions.append("(name LIKE %(search)s OR subject LIKE %(search)s OR custom_customer_name LIKE %(search)s OR customer LIKE %(search)s)")
        values["search"] = search_escaped
        
    where_clause = " AND ".join(conditions)
    
    # 1. Fetch summary counts for the logged-in technician (always calculated from all their tickets)
    summary_query = f"""
        SELECT 
            COUNT(*) as `Total`,
            SUM(CASE WHEN status = 'Open' THEN 1 ELSE 0 END) as `Open`,
            SUM(CASE WHEN status = 'Working' THEN 1 ELSE 0 END) as `Working`,
            SUM(CASE WHEN status = 'Resolved' THEN 1 ELSE 0 END) as `Resolved`,
            SUM(CASE WHEN status = 'Pending' THEN 1 ELSE 0 END) as `Pending`
        FROM `tabHD Ticket`
        WHERE JSON_SEARCH(`_assign`, 'one', %(user)s) IS NOT NULL
    """
    summary_data = frappe.db.sql(summary_query, {"user": user}, as_dict=True)
    summary = {
        "Total": 0,
        "Open": 0,
        "Working": 0,
        "Resolved": 0,
        "Pending": 0
    }
    if summary_data:
        summary = {
            "Total": int(summary_data[0].get("Total") or 0),
            "Open": int(summary_data[0].get("Open") or 0),
            "Working": int(summary_data[0].get("Working") or 0),
            "Resolved": int(summary_data[0].get("Resolved") or 0),
            "Pending": int(summary_data[0].get("Pending") or 0)
        }
        
    # 2. Fetch tickets count and ticket list
    total_count_query = f"""
        SELECT COUNT(*) 
        FROM `tabHD Ticket`
        WHERE {where_clause}
    """
    total_count = frappe.db.sql(total_count_query, values)[0][0]
    
    tickets_query = f"""
        SELECT {TICKET_COLUMNS}
        FROM `tabHD Ticket`
        WHERE {where_clause}
        ORDER BY creation DESC
        LIMIT %(limit_start)s, %(limit_page_length)s
    """

    values["limit_start"] = int(limit_start)
    values["limit_page_length"] = int(limit_page_length)

    tickets = frappe.db.sql(tickets_query, values, as_dict=True)
    tickets = enrich_tickets_customer_details(tickets)
    attach_machines_and_invoices(tickets)

    result = {
        "tickets": tickets,
        "total_count": total_count,
        "summary": summary
    }
    frappe.cache().set_value(cache_key, result, expires_in_sec=600)
    return result


# Expense Claim dialog ranges: days back from today, both ends inclusive
EXPENSE_CLAIM_RANGES = {"today": 0, "last_week": 6}

# Expense Claim Type for new lines when a request names none (pages from before the type picker)
EXPENSE_CLAIM_TYPE = "Travel"

# Those lines name their ticket first in the description: "Ticket 0627 · customer · machine"
LINE_TICKET = re.compile(r"\bTicket\s+([^\s·<]+)")


@frappe.whitelist()
def get_expense_claim_tickets(date_range: str = "today"):
    """Tickets for the Expense Claim dialog.

    A ticket is listed when the user visited it (checked in or out) within the range,
    whatever its status, which is shown as it is on the ticket. "last_week" is the last
    7 days including today.
    """
    user = frappe.session.user
    if user == "Guest":
        frappe.throw(_("Authentication required."), frappe.PermissionError)

    tickets, from_date, to_date = get_expense_claim_rows(user, date_range)
    return {"tickets": tickets, "from_date": from_date, "to_date": to_date, "expense_types": get_expense_claim_types(user)}


@frappe.whitelist()
def save_to_expense_claim(amounts: Union[dict, str], types: Union[dict, str, None] = None):
    """Store the dialog's amounts ({ticket: amount}) and Expense Claim Types ({ticket: type})
    in the user's HRMS Expense Claim.

    Each ticket is one line in the employee's open draft claim (created when needed),
    found again by the "Ticket <name>" at the start of its description; HR approves and
    submits the claim as usual. Sharing again updates the line, an empty amount removes
    it, and lines in a submitted claim can't change.
    """
    user = frappe.session.user
    if user == "Guest":
        frappe.throw(_("Authentication required."), frappe.PermissionError)

    amounts = parse_ticket_map(amounts, _("Expense amounts"))
    expense_types = parse_ticket_map(types, _("Expense Claim Types"))

    employee = frappe.db.get_value("Employee", {"user_id": user, "status": "Active"}, ["name", "company"], as_dict=True)
    if not employee:
        frappe.throw(_("No active Employee is linked to your user, so expenses can't be added to an Expense Claim."))

    tickets = {t["name"]: t for t in get_expense_claim_rows(user, "last_week")[0]}
    allowed_types = set(get_expense_claim_types(user))
    # The employee's draft claims holding ticket lines, newest first; new lines go to the newest
    drafts = [frappe.get_doc("Expense Claim", name) for name in frappe.db.sql_list("""
        SELECT claim.name FROM `tabExpense Claim` claim
        WHERE claim.employee = %s AND claim.docstatus = 0 AND EXISTS (
            SELECT 1 FROM `tabExpense Claim Detail` detail
            WHERE detail.parent = claim.name AND detail.parenttype = 'Expense Claim'
                AND detail.description LIKE %s
        )
        ORDER BY claim.creation DESC
    """, (employee.name, "%Ticket %"))]
    lines = {}
    for claim in reversed(drafts):
        for row in claim.expenses:
            line_ticket = get_line_ticket(row.description)
            if line_ticket:
                lines[line_ticket] = (claim, row)

    changed = {}
    for ticket, value in amounts.items():
        t = tickets.get(str(ticket))
        if not t:
            frappe.throw(_("Ticket {0} is not a ticket you visited in the last week.").format(ticket))
        amount = parse_expense_amount(t["name"], value)
        expense_type = expense_types.get(ticket) or None
        if expense_type and expense_type not in allowed_types:
            frappe.throw(_("{0} is not an Expense Claim Type you can use.").format(expense_type))
        if t["expense_claim_locked"]:
            if amount != t["expense_amount"] or (expense_type and expense_type != t["expense_type"]):
                frappe.throw(_("Ticket {0} is already in submitted Expense Claim {1}.").format(t["name"], t["expense_claim"]))
            continue

        if t["name"] in lines:
            claim, row = lines[t["name"]]
            if amount is None:
                claim.remove(row)
            elif flt(row.amount) != amount or (expense_type and expense_type != row.expense_type):
                row.amount = row.sanctioned_amount = amount
                row.expense_date = getdate(t["expense_date"])
                if expense_type and expense_type != row.expense_type:
                    row.expense_type = expense_type
                    row.default_account = None  # HRMS fetches the new type's account on save
            else:
                continue
        elif amount is not None:
            if not expense_type:
                if types is not None:
                    frappe.throw(_("Select an Expense Claim Type for ticket {0}.").format(t["name"]))
                expense_type = EXPENSE_CLAIM_TYPE
            if not drafts:
                drafts.append(new_expense_claim(employee))
            claim = drafts[0]
            claim.append("expenses", {
                "expense_date": getdate(t["expense_date"]),
                "expense_type": expense_type,
                # "Ticket <name>" stays untranslated: it is how the line is found again
                "description": " · ".join(filter(None, (f"Ticket {t['name']}", t.get("custom_customer_name"), t["machine_display"]))),
                "amount": amount,
                "sanctioned_amount": amount,
                "cost_center": claim.cost_center,
            })
        else:
            continue
        changed[id(claim)] = claim

    for claim in changed.values():
        if claim.expenses:
            claim.save(ignore_permissions=True)
        elif not claim.is_new():
            frappe.delete_doc("Expense Claim", claim.name, ignore_permissions=True)

    open_claim = next((c.name for c in drafts if c.expenses and not c.is_new()), None)
    return {"claim": open_claim, "changed": len(changed)}


@frappe.whitelist()
def download_expense_claim(
    date_range: str = "today", amounts: Union[dict, str, None] = None, types: Union[dict, str, None] = None
):
    """The Expense Claim dialog's Share (Excel) file: the same tickets with the dialog's
    amounts ({ticket: amount}) and Expense Claim Types ({ticket: type}) and a total row,
    as a real .xlsx."""
    user = frappe.session.user
    if user == "Guest":
        frappe.throw(_("Authentication required."), frappe.PermissionError)

    amounts = parse_ticket_map(amounts, _("Expense amounts"))
    types = parse_ticket_map(types, _("Expense Claim Types"))

    tickets, from_date, to_date = get_expense_claim_rows(user, date_range)
    typed = {t["name"]: parse_expense_amount(t["name"], amounts.get(t["name"])) for t in tickets}
    expense_types = {t["name"]: types.get(t["name"]) or t["expense_type"] for t in tickets if typed[t["name"]] is not None}
    frappe.response["filename"] = f"expense_claim_{date_range}_{today()}.xlsx"
    frappe.response["filecontent"] = build_expense_claim_xlsx(user, tickets, from_date, to_date, typed, expense_types)
    frappe.response["type"] = "binary"


def get_expense_claim_rows(user, date_range):
    """(tickets, from_date, to_date) shown in the Expense Claim dialog and its Excel file,
    each with the amount already in the user's Expense Claim, if any."""
    from_date, to_date = get_expense_claim_range(date_range)
    tickets = get_claimable_tickets(user, from_date, to_date, TICKET_COLUMNS)
    tickets = enrich_tickets_customer_details(tickets)
    attach_machines_and_invoices(tickets)
    employee = frappe.db.get_value("Employee", {"user_id": user, "status": "Active"}, "name")
    claimed = get_claimed_ticket_lines(employee, [t["name"] for t in tickets])
    for t in tickets:
        machines = (m.get("machine_name") or m.get("machine_type") for m in t["custom_machine_type_list"])
        t["machine_display"] = ", ".join(filter(None, machines)) or t.get("custom_machine_name") or ""
        t["location"] = ", ".join(filter(None, (t.get("custom_address"), t.get("custom_city__district_"))))
        line = claimed.get(t["name"])
        t["expense_amount"] = line.amount if line else None
        t["expense_type"] = line.expense_type if line else None
        t["expense_claim"] = line.claim if line else None
        t["expense_claim_locked"] = bool(line and line.docstatus == 1)
    return tickets, from_date, to_date


def get_claimed_ticket_lines(employee, tickets):
    """{ticket: line} for the employee's non-cancelled Expense Claim lines of these tickets."""
    if not employee or not tickets:
        return {}
    lines = frappe.db.sql("""
        SELECT detail.description, detail.amount, detail.expense_type, claim.name AS claim, claim.docstatus
        FROM `tabExpense Claim Detail` detail
        JOIN `tabExpense Claim` claim ON claim.name = detail.parent AND detail.parenttype = 'Expense Claim'
        WHERE claim.employee = %(employee)s AND claim.docstatus < 2 AND detail.description LIKE %(marker)s
        ORDER BY claim.creation
    """, {"employee": employee, "marker": "%Ticket %"}, as_dict=True)
    wanted = set(tickets)
    claimed = {}
    for line in lines:
        line_ticket = get_line_ticket(line.description)
        if line_ticket in wanted:
            claimed[line_ticket] = line  # a later claim wins
    return claimed


def get_line_ticket(description):
    """The ticket an Expense Claim line was added for, read from its description."""
    match = LINE_TICKET.search(frappe.utils.strip_html(description or ""))
    return match.group(1) if match else None


def get_expense_claim_types(user):
    """Expense Claim Types the user can claim: those with a default account for their company."""
    company = frappe.db.get_value("Employee", {"user_id": user, "status": "Active"}, "company")
    return sorted(set(frappe.get_all("Expense Claim Account", filters={"company": company} if company else {}, pluck="parent")))


def parse_ticket_map(value, label):
    """A {ticket: value} map sent by the dialog (JSON string or dict); empty when not sent."""
    if isinstance(value, str):
        value = frappe.parse_json(value)
    if not isinstance(value or {}, dict):
        frappe.throw(_("{0} must be a map of ticket to value.").format(label))
    return value or {}


def new_expense_claim(employee):
    """An empty draft Expense Claim for the employee, filled in as the HRMS form does."""
    import erpnext
    from hrms.hr.doctype.expense_claim.expense_claim import get_default_payable_account

    return frappe.get_doc({
        "doctype": "Expense Claim",
        "employee": employee.name,
        "company": employee.company,
        "posting_date": today(),
        "expense_approver": frappe.db.get_value("Employee", employee.name, "expense_approver"),
        "payable_account": get_default_payable_account(employee.company),
        "cost_center": erpnext.get_default_cost_center(employee.company),
        "remark": _("Ticket expenses from the Technician Portal"),
    })


# Share (Excel) sheet columns: header, width
EXPENSE_CLAIM_COLUMNS = (
    ("Ticket No", 12), ("Machine", 34), ("Customer", 24), ("Required Spares", 18),
    ("Location", 42), ("Date", 15), ("Status", 12), ("Expense Type", 15), ("Expense Amount", 18),
)


def build_expense_claim_xlsx(user, tickets, from_date, to_date, amounts, expense_types):
    """Title, header, one row per ticket and a total row, all in 13pt (the requested
    default font size; the columns default to it too, for anything typed later)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    font, bold = Font(size=13), Font(size=13, bold=True)
    period = formatdate(from_date) if from_date == to_date else f"{formatdate(from_date)} – {formatdate(to_date)}"

    wb = Workbook()
    ws = wb.active
    ws.title = _("Expense Claim")
    ws.append([f"{_('Expense Claim')}: {get_fullname(user)} ({period})"])
    ws.append([_(label) for label, width in EXPENSE_CLAIM_COLUMNS])
    for t in tickets:
        ws.append([
            t["name"], t["machine_display"], t.get("custom_customer_name") or "", t.get("custom_required_spares") or "-",
            t["location"], getdate(t["expense_date"]), t.get("status"), expense_types.get(t["name"]), amounts.get(t["name"]),
        ])
    ws.append([None] * 7 + [_("Total"), sum(flt(amounts.get(t["name"])) for t in tickets)])

    total_row = ws.max_row
    for row in ws.iter_rows():
        for cell in row:
            cell.font = bold if cell.row in (1, 2, total_row) else font
            if 2 < cell.row < total_row:
                cell.alignment = Alignment(vertical="top", wrap_text=cell.column in (2, 5))
    for row in range(3, total_row):
        ws.cell(row, 1).number_format = "@"  # keep ticket numbers like 0627 as text
        ws.cell(row, 6).number_format = get_user_date_format()
        ws.cell(row, 9).number_format = "#,##0.00"
    ws.cell(total_row, 8).alignment = Alignment(horizontal="right")
    ws.cell(total_row, 9).number_format = "#,##0.00"

    for idx, (label, width) in enumerate(EXPENSE_CLAIM_COLUMNS, start=1):
        column = ws.column_dimensions[get_column_letter(idx)]
        column.width = width
        column.font = font
    ws.freeze_panes = "A3"

    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def get_expense_claim_range(date_range):
    """(from_date, to_date) for an Expense Claim range key."""
    if date_range not in EXPENSE_CLAIM_RANGES:
        frappe.throw(_("Unknown date range: {0}").format(date_range))
    to_date = getdate(today())
    return add_days(to_date, -EXPENSE_CLAIM_RANGES[date_range]), to_date


def get_claimable_tickets(user, from_date, to_date, columns):
    """Tickets the user visited (a check-in or check-out at the customer) between the dates
    (inclusive), whatever their status. `expense_date` is the user's last visit."""
    visits = """
        FROM `tabHD Ticket Check Log` check_log
        WHERE check_log.parenttype = 'HD Ticket'
            AND check_log.parent = `tabHD Ticket`.name
            AND check_log.technician = %(user)s
    """
    return frappe.db.sql(f"""
        SELECT {columns}, (SELECT MAX(check_log.timestamp) {visits}) AS expense_date
        FROM `tabHD Ticket`
        WHERE EXISTS (SELECT 1 {visits} AND DATE(check_log.timestamp) BETWEEN %(from_date)s AND %(to_date)s)
        ORDER BY expense_date DESC
    """, {"user": user, "from_date": from_date, "to_date": to_date}, as_dict=True)


def parse_expense_amount(ticket, value):
    """The typed amount as a number, or None when left empty."""
    if value is None or str(value).strip() == "":
        return None
    try:
        amount = float(value)
    except (TypeError, ValueError):
        amount = math.nan
    if not math.isfinite(amount) or amount < 0:
        frappe.throw(_("Enter a valid, non-negative expense amount for ticket {0}.").format(ticket))
    return amount


def attach_machines_and_invoices(tickets):
    """Add each ticket's machine rows and latest non-cancelled Sales Invoice."""
    if not tickets:
        return

    ticket_names = [t["name"] for t in tickets]
    child_rows = frappe.get_all(
        "Machine type list",
        filters={"parent": ["in", ticket_names], "parenttype": "HD Ticket"},
        fields=["parent", "machine_type", "machine_name", "machine_brand", "machine_quantity", "machine_problem", "purchased_at_scs", "purchase_year"],
        order_by="idx asc"
    )
    child_map = {}
    for row in child_rows:
        p = row["parent"]
        if p not in child_map:
            child_map[p] = []
        child_map[p].append(row)

    # Fetch any linked Sales Invoices
    inv_rows = frappe.get_all(
        "Sales Invoice",
        filters={"custom_hd_ticket": ["in", ticket_names], "docstatus": ["!=", 2]},
        fields=["name", "custom_hd_ticket", "grand_total", "docstatus"],
        order_by="creation desc"
    )
    inv_map = {}
    for inv in inv_rows:
        tk = inv.get("custom_hd_ticket")
        if tk and tk not in inv_map:
            inv_map[tk] = inv

    for t in tickets:
        t["custom_machine_type_list"] = child_map.get(t["name"], [])
        t["sales_invoice"] = inv_map.get(t["name"])
